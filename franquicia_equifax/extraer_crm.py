"""Extrae la planilla de gestión de prospectos (una fila por prospecto, p. ej. "PROSPECCIÓN MARTI MDZ").

La planilla tiene columnas de estado (producto, probabilidad, "Contactada por…"), pero casi toda la
información real está en el texto libre de "Comentarios". Por eso cada prospecto se clasifica leyendo
ese texto con reglas simples y explicables (ver ETAPAS). Teléfonos y mails NO se copian: se borran del
texto antes de guardarlo.

Uso:
    python extraer_crm.py RUTA.xlsx [salida.json] [--anio 2026]
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

# Etapa final de cada prospecto. Se evalúan en orden: la primera que coincide gana.
ETAPAS = [
    ("Ya era cliente", r"ya es cl[ie]+nte|ya utiliza equifax|comunicad[ao] con mariano"),
    ("Alta lograda", r"gener[oó] el contrato|contrato y es client"),
    ("Número inválido", r"no corresponde a un abonado|caracter[ií]stica solicitada|me dio equivocado|no est[aá] disponible\" y luego"),
    ("Cierre pendiente", r"esperando que (nos )?envi[eé] (los datos|lo de)|realizar el contrato|generar el contrato"),
    ("Demo o reunión", r"(?<!no quiere )\bdemos?\b|videollamada|\bmeet\b|la reuni[oó]n se realiz|realizar el contrato|lo de atm|personalmente a la oficina|presencialmente"),
    ("Contacto clave sin llamar", r"interesante para prospectar|especial inter[eé]s para prospectar"),
    ("No interesado", r"no est[aá] interesad|no esta interesad|no le interesa|no quiere (incurrir|un gasto|gastar)|jubil|retirad|fallec|me cort[oó]|no est[aá] trabajando|no administra|dej[oó] de trabajar|se dedica (m[aá]s que nada )?a la (venta|construcci)"),
    ("Interesado en seguimiento", r"interesad|pack de|info por (wpp|app|awpp|mail)|(envi|mand)[eé] info|hacer seguimiento|ro le mand"),
    ("Pidió que lo llamen", r"me pidi[oó] que (lo|la) (vuelva a )?(llam|contact)|vuelva a llamar|ya me devolv|va a enviarme|dsp de las"),
    ("No contesta", r"no contest|no atendi|no atiende|ocupado|buz[oó]n|apagado|no me contest"),
]
COMPETIDORES = {
    "Codeme": r"codeme",
    "Nosis": r"nos+is",
    "Riesgo Online": r"rie[sg]+o online",
    "Banco Central": r"banco central",
    "Info Experto": r"info exp",
    "Treni": r"treni",
    "Bolsa de Comercio": r"bolsa de comercio",
}
OBJECIONES = {
    "Costo fijo / prefiere prepago": r"costo fijo|gasto fijo|gasto mensual|prepago|tiene un costo|mejor precio",
    "Ya no trabaja o se retira": r"jubil|retirad|no est[aá] trabajando|dej[oó] de trabajar|no administra|no est[aá] ejerciendo",
    "Lo decide otra persona": r"socia|socio|hija|hijo|esposa|hermana|secretaria|sobrin",
    "Desconfianza del canal": r"desconfiad|no va a dar ning|no me cre[ií]a",
}
PRODUCTOS = {
    "Pack de reportes": r"pack|reportes",
    "Mora Control": r"mora control",
    "Porfolio Monitor": r"porfolio",
}
FRANJAS = {"Mañana": r"mañana en la mañana|en la mañana|antes de las 13|media mañana|11(:30)? hs|10:30|de la mañana",
           "Siesta / mediodía": r"siesta|13:30|medio d[ií]a",
           "Tarde": r"en la tarde|16:30|16 hs|18 hs|19 hs|dsp de las 19|antes de las 18"}
SEGUIMIENTO = r"volv[ií] a (llamar|contactar|intentar|escribir)|volver a (contactar|llamar|intentar|probar)|llamar (lunes|martes|mi[eé]rcoles|jueves|viernes)|hacer seguimiento|seguir insistiendo"
WHATSAPP = r"wpp|whats|awpp|\bapp\b"


def norm(s: str) -> str:
    return unicodedata.normalize("NFC", str(s or "")).strip()


def limpiar(txt: str) -> str:
    """Borra mails y números de teléfono del comentario."""
    txt = re.sub(r"\S+@\S+", "[mail]", txt)
    return re.sub(r"\+?\d[\d\s-]{6,}\d", "[tel]", txt)


def clasificar(texto: str) -> str:
    t = texto.lower()
    for etapa, patron in ETAPAS:
        if re.search(patron, t):
            return etapa
    return "Otro"


def etiquetas(texto: str, tabla: dict) -> list[str]:
    t = texto.lower()
    return [k for k, p in tabla.items() if re.search(p, t)]


def fecha(v, anio: int):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    m = re.fullmatch(r"\s*(\d{1,2})\s*[-/]\s*(\d{1,2})\s*", str(v or ""))
    return dt.date(anio, int(m.group(2)), int(m.group(1))).isoformat() if m else None


def main(xlsx: str, salida: str, anio: int) -> None:
    ws = openpyxl.load_workbook(xlsx, data_only=True).worksheets[0]
    cab = [norm(c).lower() for c in next(ws.iter_rows(max_row=1, values_only=True))]
    col = lambda *claves: next((i for i, c in enumerate(cab) if any(k in c for k in claves)), None)
    c_id, c_nom, c_fecha, c_com = col("id"), col("nombre"), col("fecha"), col("comentario")
    c_prod, c_prob = col("producto"), col("probabilidad")
    c_ro = col("rocio")

    prospectos, ultima_fecha = [], None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r):
            continue
        com = norm(r[c_com]) if c_com is not None else ""
        if com in ("", "-"):
            com = ""
        f = fecha(r[c_fecha], anio) if c_fecha is not None else None
        if f:
            ultima_fecha = f
        base_heredada = bool(re.search(r"ex contacto de rami|recontactado por ro", com.lower()))
        prospectos.append({
            "id": int(r[c_id]) if isinstance(r[c_id], (int, float)) else None,
            "nombre": norm(r[c_nom]).title() if c_nom is not None else "",
            "fecha": f,
            "fecha_aprox": f or (ultima_fecha if com else None),
            "gestionado": bool(com),
            "etapa": clasificar(com) if com else "Sin gestionar",
            "producto_planilla": norm(r[c_prod]) if c_prod is not None else "",
            "probabilidad": norm(r[c_prob]) if c_prob is not None else "",
            "competidores": etiquetas(com, COMPETIDORES),
            "objeciones": etiquetas(com, OBJECIONES),
            "productos": etiquetas(com, PRODUCTOS),
            "franjas": etiquetas(com, FRANJAS),
            "seguimientos": len(re.findall(SEGUIMIENTO, com.lower())),
            "whatsapp": bool(re.search(WHATSAPP, com.lower())),
            "rocio": base_heredada or (c_ro is not None and "sin contactar" not in norm(r[c_ro]).lower()),
            "base_heredada": base_heredada,
            "comentario": limpiar(com),
        })
    Path(salida).parent.mkdir(parents=True, exist_ok=True)
    Path(salida).write_text(json.dumps({"fuente": Path(xlsx).name, "anio": anio, "prospectos": prospectos},
                                       ensure_ascii=False, indent=1))
    from collections import Counter
    print(f"{len(prospectos)} prospectos -> {salida}")
    print("  " + ", ".join(f"{k}: {v}" for k, v in Counter(p["etapa"] for p in prospectos).most_common()))


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    anio = int(sys.argv[sys.argv.index("--anio") + 1]) if "--anio" in sys.argv else 2026
    args = [a for a in args if a != str(anio)]
    if not args:
        sys.exit(__doc__)
    main(args[0], args[1] if len(args) > 1 else str(Path(__file__).parent / "datos" / "crm.json"), anio)
