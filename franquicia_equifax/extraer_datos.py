"""Extrae el Excel de liquidaciones mensuales de Virtus a una serie limpia (JSON).

Cada hoja del Excel es una liquidación: lo que Equifax transfirió a Virtus ese
mes, los gastos que se descontaron, el fondo de resguardo y el reparto entre
los 4 socios. Este script normaliza todas las hojas (que cambiaron de formato
con el tiempo) a un mismo esquema y clasifica cada gasto en un rubro.

Uso:
    python extraer_datos.py RUTA_AL_EXCEL.xlsx [salida.json]

Los datos de salida NO se versionan (ver .gitignore): contienen información
financiera privada de la sociedad.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

SOCIOS = ["JPM", "GF", "MG", "MC"]

# Hoja -> (mes de liquidación AAAA-MM, mes de ventas que se cobra, nota)
HOJAS = {
    "Mayo (abril)": ("2025-05", "2025-04", "Cifras idénticas al bloque 'Abril' de la hoja 'Meses anteriores': revisar si es duplicado."),
    "Junio (mayo)": ("2025-06", "2025-05", ""),
    "Julio (junio)": ("2025-07", "2025-06", ""),
    "Agosto (julio)": ("2025-08", "2025-07", ""),
    "Septiembre (agosto)": ("2025-09", "2025-08", "Equifax retuvo el IVA de la factura (21%)."),
    "Octubre (septiembre)": ("2025-10", "2025-09", ""),
    "Noviembre (": ("2025-11", "2025-10", ""),
    "Diciembre (novimembre y cierre)": ("2025-12", "2025-11", ""),
    "Enero parcial (diciembre)": ("2026-01", "2025-12", "Diciembre se cobró en dos partes; se suman."),
    "Enero parcial 2 (diciembre - en": ("2026-01", "2025-12", "Diciembre se cobró en dos partes; se suman."),
    "Febrero ( enero, en marzo)": ("2026-02", "2026-01", "Las ventas de enero entraron recién en marzo."),
    "Marzo (febrero)": ("2026-03", "2026-02", ""),
    "Abril (marzo)": ("2026-04", "2026-03", "Se postergaron ~2,8 M de gastos a la liquidación siguiente."),
    "Mayo(abril)": ("2026-05", "2026-04", ""),
    "Junio (mayo26)": ("2026-06", "2026-05", ""),
    "Julio (junio26)": ("2026-07", "2026-06", ""),
    "Agosto (julio - entró en sept)": ("2026-08", "2026-07", "Las ventas de julio entraron en septiembre."),
    "Septiembre (agosto + sobrantes ": ("2026-09", "2026-08", "Primer mes con variable comercial y nuevos porcentajes de reparto."),
}

# Rubros: el orden importa (el primer patrón que coincide gana).
RUBROS = [
    ("Fee Equifax", r"fee virtus"),
    ("Bancos e imp. al cheque", r"cr[eé]d.*d[eé]b|sicreb|mantenimiento banco"),
    ("Equipo comercial", r"comercial|rocio|martina|lucila|sindicato|gustavo reyes|vep atrasado|variable ro"),
    ("Impuestos", r"iva|ganancia|ganacia|vep|iibb|impuesto|sellos"),
    ("Call center", r"call center|llamadas"),
    ("Contadora", r"contadora"),
    ("Oficina y alquiler", r"alquiler|tel[eé]fono"),
    ("Tarjeta de crédito", r"^tarjeta$"),
    ("Promoción y merchandising", r"amigas|altas|evento|sponsor|sticker|mates|impresiones|tarjetas personales"),
    ("Reintegros y viajes", r"reintegr|viaje|gastos ro|mariano gastos|juan pablo \(lucio\)|lucio"),
]
OTROS = "Otros / extraordinarios"


def rubro(label: str) -> str:
    t = label.strip().lower()
    for nombre, patron in RUBROS:
        if re.search(patron, t):
            return nombre
    return OTROS


def num(v) -> float | None:
    """Convierte celdas numéricas o textos tipo '1.740.000' a float."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):
        return float(s.replace(".", "").replace(",", "."))
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def miles(v) -> float:
    """En 'Meses anteriores' los montos chicos vienen como 812.0 = 812.000."""
    x = num(v) or 0.0
    return x * 1000 if 0 < abs(x) < 10000 else x


def parse_hoja(ws) -> dict:
    get = lambda r, c: ws.cell(r, c).value
    disponible = num(get(5, 1))

    # Cabecera de socios (fila 3) con % (fila 4) y montos base (fila 5).
    shares, base = {}, {}
    for c in range(2, 10):
        h = str(get(3, c) or "").strip().upper()
        if h in SOCIOS:
            shares[h] = num(get(4, c))
            base[h] = num(get(5, c))

    pagado = saldo = factura = None
    inicio = None
    for r in range(6, 13):
        lab = str(get(r, 1) or "").lower()
        v = num(get(r, 2))
        if "pagado virtus" in lab:
            pagado = v
        elif any(k in lab for k in ("quedó en cuenta", "a favor", "saldo mes")):
            saldo = v
        elif "factua" in lab:
            factura = v
        if v is not None and disponible and abs(v - disponible) < 1:
            inicio = r  # la fila "total disponible" es la última que coincide
    if pagado is None:
        pagado, saldo = disponible, 0.0

    gastos, resguardo, total_reportado, diferidos = [], 0.0, None, 0.0
    for r in range(inicio + 1, inicio + 30):
        lab = str(get(r, 1) or "").strip()
        v = num(get(r, 2))
        low = lab.lower()
        if low.startswith("total gastos"):
            mas_fondo = "más" in low or "mas " in low
            if total_reportado is None and v is not None:
                total_reportado = v - resguardo if mas_fondo else v
            if mas_fondo:
                break
            continue
        if "resguardo" in low:
            resguardo = v or 0.0
            if total_reportado is not None:
                break
            continue
        if low.startswith("neto") or "transferencia" in low or "base proporcional" in low:
            break
        nota_c = str(get(r, 3) or "").lower()
        if "pr" in nota_c and "xima" in nota_c and num(get(r, 4)):
            diferidos += num(get(r, 4))
        if lab and v:
            gastos.append({"concepto": lab, "rubro": rubro(lab), "monto": v})

    # Variable comercial (desde ago-2026): C13 = "Variable Ro", D13 = monto.
    if "variable" in str(get(13, 3) or "").lower() and num(get(13, 4)):
        gastos.append({"concepto": "Variable comercial (Rocío)", "rubro": "Equipo comercial", "monto": num(get(13, 4))})

    suma = sum(g["monto"] for g in gastos if g["concepto"] != "Variable comercial (Rocío)")
    if total_reportado and abs(total_reportado - suma) > 1:
        gastos.append({"concepto": "Diferencia sin detalle en la hoja", "rubro": OTROS, "monto": total_reportado - suma})

    # Reintegros a socios (bloque "Transferencias socios final").
    reintegros = {s: 0.0 for s in SOCIOS}
    for r in range(25, 60):
        if str(get(r, 1) or "").strip().lower() == "gastos reintegrables":
            for c in range(2, 8):
                h = str(get(r - 1, c) or "").upper()
                cod = ("JPM" if "JPM" in h else "GF" if h.strip() == "GF" else
                       "MG" if ("MARIO" in h and "MARIANO" not in h) or h.strip().startswith("MG") else
                       "MC" if "MARIANO" in h or h.strip().startswith("MC") else None)
                if cod and num(get(r, c)):
                    reintegros[cod] += num(get(r, c))
            break

    return {
        "ingreso": pagado or 0.0,
        "saldo_anterior": saldo or 0.0,
        "facturado_con_iva": factura,
        "disponible": disponible,
        "gastos": gastos,
        "resguardo": resguardo,
        "shares": shares,
        "reparto": base,
        "reintegros": reintegros,
        "diferidos": diferidos,
    }


def meses_anteriores(ws) -> list[dict]:
    """Bloques manuales de ene-abr 2025 (formato texto, en miles)."""
    g = lambda ref: ws[ref].value
    bloques = [
        ("2025-01", "2024-12", "D20", [("Impuestos", "D21"), ("Comercial (10%)", "D22"), ("Contadora", "D23"), ("Call center", "D24")], "D26", (0.4, 0.3, 0.15, 0.15)),
        ("2025-02", "2025-01", "D32", [("Impuestos", "D33"), ("Comercial (10%)", "D34"), ("Contadora", "D35"), ("Call center", "D36")], "D37", (0.4, 0.3, 0.15, 0.15)),
        ("2025-03", "2025-02", "D43", [("Impuestos", "D44"), ("Comercial (10%)", "D45"), ("Contadora", "D46"), ("Call center", "D47")], "D48", (0.4, 0.3, 0.15, 0.15)),
        ("2025-04", "2025-03", "D55", [("Impuestos", "D56"), ("Comercial FG", "D57"), ("Contadora", "D58"), ("Call center", "D59"), ("Gastos", "D60")], "D61", (0.4, 0.3, 0.15, 0.15)),
    ]
    out = []
    for cobro, venta, ing, items, rep, sh in bloques:
        reparto_total = miles(g(rep))
        out.append({
            "hoja": "Meses anteriores", "mes": cobro, "venta": venta,
            "nota": "Formato resumido de 2025: sin detalle de tarjeta, alquiler ni bancos.",
            "ingreso": miles(g(ing)), "saldo_anterior": 0.0, "facturado_con_iva": None,
            "disponible": miles(g(ing)),
            "gastos": [{"concepto": c, "rubro": rubro(c) if c != "Gastos" else OTROS, "monto": miles(g(ref))} for c, ref in items],
            "resguardo": 0.0,
            "shares": dict(zip(SOCIOS, sh)),
            "reparto": {s: reparto_total * p for s, p in zip(SOCIOS, sh)},
            "reintegros": {s: 0.0 for s in SOCIOS}, "diferidos": 0.0,
        })
    # Nov-24 y dic-24: solo se conoce el monto repartido.
    for cobro, ref in (("2024-11", "C6"), ("2024-12", "C10")):
        t = miles(g(ref))
        out.append({"hoja": "Meses anteriores", "mes": cobro, "venta": None, "solo_reparto": True,
                    "nota": "Solo se registró el monto repartido.", "ingreso": None, "saldo_anterior": 0.0,
                    "facturado_con_iva": None, "disponible": None, "gastos": [], "resguardo": 0.0,
                    "shares": dict(zip(SOCIOS, (0.4, 0.3, 0.15, 0.15))),
                    "reparto": {s: t * p for s, p in zip(SOCIOS, (0.4, 0.3, 0.15, 0.15))},
                    "reintegros": {s: 0.0 for s in SOCIOS}, "diferidos": 0.0})
    return out


def mayo_2025(ws) -> dict:
    """La hoja 'Mayo (abril)' de 2025 usa el formato viejo en miles."""
    g = lambda ref: ws[ref].value
    items = [("Impuestos", "B8"), ("Comercial FG", "B9"), ("Contadora", "B10"), ("Call Center", "B11"), ("Gastos", "B12")]
    sh = (0.4, 0.3, 0.15, 0.15)
    rep = miles(g("B13"))
    return {"ingreso": miles(g("B7")), "saldo_anterior": 0.0, "facturado_con_iva": None, "disponible": miles(g("B7")),
            "gastos": [{"concepto": c, "rubro": rubro(c) if c != "Gastos" else OTROS, "monto": miles(g(r))} for c, r in items],
            "resguardo": 0.0, "shares": dict(zip(SOCIOS, sh)), "reparto": {s: rep * p for s, p in zip(SOCIOS, sh)},
            "reintegros": {"JPM": num(g("I4")) or 0.0, "GF": 0.0, "MG": 0.0, "MC": num(g("J4")) or 0.0}, "diferidos": 0.0}


def fusionar(a: dict, b: dict) -> dict:
    """Suma dos liquidaciones parciales del mismo mes."""
    out = dict(a)
    for k in ("ingreso", "resguardo", "diferidos"):
        out[k] = (a[k] or 0) + (b[k] or 0)
    out["disponible"] = a["disponible"] + b["disponible"]
    out["saldo_anterior"] = a["saldo_anterior"] + b["saldo_anterior"]
    out["gastos"] = a["gastos"] + b["gastos"]
    out["reparto"] = {s: a["reparto"].get(s, 0) + b["reparto"].get(s, 0) for s in SOCIOS}
    out["reintegros"] = {s: a["reintegros"][s] + b["reintegros"][s] for s in SOCIOS}
    out["hoja"] = a["hoja"] + " + " + b["hoja"]
    return out


def pagado_a_socios(ws) -> dict:
    """Control cruzado: hoja 'Pagado a socios' (ene-ago 2025)."""
    meses = {7 + i: f"2025-{i:02d}" for i in range(1, 9)}  # col C..J
    cod = {"Mariano Comba": "MC", "Mario Gaspar": "MG", "Juan Pablo": "JPM", "Gustavo": "GF"}
    out = {}
    for r in range(8, 12):
        socio = cod.get(str(ws.cell(r, 2).value).strip())
        for c in range(3, 11):
            v = ws.cell(r, c).value
            if socio and v is not None:
                out.setdefault(meses[c + 5], {})[socio] = num(v) if not isinstance(v, str) else num(v)
    return out


def main(xlsx: str, salida: str) -> None:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    registros: dict[str, dict] = {}
    for hoja, (mes, venta, nota) in HOJAS.items():
        ws = wb[hoja]
        d = mayo_2025(ws) if hoja == "Mayo (abril)" else parse_hoja(ws)
        d.update({"hoja": hoja, "mes": mes, "venta": venta, "nota": nota})
        registros[mes] = fusionar(registros[mes], d) if mes in registros else d
    for d in meses_anteriores(wb["Meses anteriores"]):
        registros.setdefault(d["mes"], d)

    serie = [registros[k] for k in sorted(registros)]
    for d in serie:
        d["total_gastos"] = sum(g["monto"] for g in d["gastos"])
        d["neto_socios"] = sum(d["reparto"].values())
    Path(salida).write_text(json.dumps({
        "fuente": Path(xlsx).name,
        "socios": SOCIOS,
        "rubros": [r for r, _ in RUBROS] + [OTROS],
        "meses": serie,
        "control_pagado_a_socios": pagado_a_socios(wb["Pagado a socios"]),
    }, ensure_ascii=False, indent=1))
    print(f"{len(serie)} meses -> {salida}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else str(Path(__file__).parent / "datos" / "virtus_mensual.json"))
