"""Extrae el Excel de prospección telefónica (llamadas por comercial y por día).

Cada hoja es un mes. Dentro de la hoja, cada comercial tiene un bloque: una fila
con su nombre y las fechas, y debajo las filas "Llamadas realizadas",
"Prospectos que respondieron", "Interesados", "No interesados" y los totales
"CERRADO" / "Por cerrar" del mes.

Uso:
    python extraer_prospeccion.py RUTA_AL_EXCEL.xlsx [salida.json] [--anio 2026]

El Excel no registra la hora de cada llamada: solo se puede analizar por día.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6, "JULIO": 7,
         "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12}
METRICAS = [("no interesad", "no_interesados"), ("interesad", "interesados"),
            ("respondieron", "respondieron"), ("llamadas", "llamadas")]


def fecha(txt, anio: int) -> dt.date | None:
    m = re.fullmatch(r"\s*(\d{1,2})\s*[-/]+\s*(\d{1,2})\s*", str(txt or ""))
    if not m:
        return None
    try:
        return dt.date(anio, int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def primer_numero(ws, r: int, c: int):
    for cc in range(c + 1, ws.max_column + 1):
        v = ws.cell(r, cc).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return float(v)
    return None


def main(xlsx: str, salida: str, anio: int) -> None:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    dias: dict[tuple, dict] = {}
    cierres, avisos = [], []
    # En orden cronológico: así una hoja copiada de la anterior se detecta como repetida.
    hojas = sorted(wb.worksheets, key=lambda w: MESES.get(w.title.strip().upper(), 99))
    for ws in hojas:
        mes = MESES.get(ws.title.strip().upper())
        # Filas de encabezado: nombre en la columna B y fechas a la derecha.
        cabeceras = [r for r in range(1, ws.max_row + 1)
                     if str(ws.cell(r, 2).value or "").isupper() and any(fecha(ws.cell(r, c).value, anio) for c in range(3, 16))]
        for i, r0 in enumerate(cabeceras):
            comercial = str(ws.cell(r0, 2).value).strip().title()
            r1 = cabeceras[i + 1] if i + 1 < len(cabeceras) else ws.max_row + 1
            cols = {c: fecha(ws.cell(r0, c).value, anio) for c in range(3, 16)}
            filas = {}
            for r in range(r0 + 1, r1):
                lab = str(ws.cell(r, 2).value or "").lower()
                for patron, clave in METRICAS:
                    if patron in lab:
                        filas[clave] = r
                        break
            nuevos = 0
            for c, f in cols.items():
                if not f or ws.cell(filas["llamadas"], c).value in (None, ""):
                    continue
                reg = {k: float(ws.cell(r, c).value or 0) for k, r in filas.items()}
                key = (comercial, f.isoformat())
                if key in dias:
                    if dias[key]["valores"] != reg:
                        avisos.append(f"{comercial} {f:%d/%m}: valores distintos en {dias[key]['hoja']} y {ws.title}; se usa {dias[key]['hoja']}.")
                    continue
                if mes and f.month != mes:
                    avisos.append(f"{comercial} {f:%d/%m} figura en la hoja {ws.title}.")
                dias[key] = {"hoja": ws.title, "valores": reg}
                nuevos += 1
            if nuevos == 0:
                avisos.append(f"Bloque de {comercial} en {ws.title} repetido de otra hoja: se ignora.")
                continue
            # Totales del mes: CERRADO y Por cerrar (texto en cualquier columna).
            cerr = {"comercial": comercial, "hoja": ws.title, "mes": mes, "cerrados": 0.0, "por_cerrar": 0.0}
            for r in range(r0 + 1, r1):
                for c in range(2, ws.max_column + 1):
                    t = str(ws.cell(r, c).value or "").lower().replace('"', "").strip()
                    if t == "cerrado":
                        cerr["cerrados"] = primer_numero(ws, r, c) or 0.0
                    elif t == "por cerrar":
                        cerr["por_cerrar"] = primer_numero(ws, r, c) or 0.0
            cierres.append(cerr)

    registros = []
    for (comercial, f), d in sorted(dias.items(), key=lambda x: (x[0][1], x[0][0])):
        v = d["valores"]
        registros.append({"comercial": comercial, "fecha": f, "hoja": d["hoja"],
                          "llamadas": v.get("llamadas", 0), "respondieron": v.get("respondieron", 0),
                          "interesados": v.get("interesados", 0), "no_interesados": v.get("no_interesados", 0)})
    Path(salida).parent.mkdir(parents=True, exist_ok=True)
    Path(salida).write_text(json.dumps({"fuente": Path(xlsx).name, "anio": anio, "dias": registros,
                                        "cierres": cierres, "avisos": avisos}, ensure_ascii=False, indent=1))
    print(f"{len(registros)} días, {len(cierres)} bloques de cierre -> {salida}")
    for a in avisos:
        print("  aviso:", a)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    anio = int(sys.argv[sys.argv.index("--anio") + 1]) if "--anio" in sys.argv else 2026
    args = [a for a in args if a != str(anio)]
    if not args:
        sys.exit(__doc__)
    main(args[0], args[1] if len(args) > 1 else str(Path(__file__).parent / "datos" / "prospeccion.json"), anio)
